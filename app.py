from flask import Flask, request, send_file, jsonify
from werkzeug.utils import secure_filename
from pathlib import Path
from PIL import Image
from io import BytesIO
import fitz, pikepdf, os, tempfile, subprocess, zipfile, csv, json, re, shutil
from docx import Document
from docx.shared import Inches
from openpyxl import load_workbook, Workbook
from pptx import Presentation
from pptx.util import Inches as PptInches
import pytesseract

app = Flask(__name__, static_folder='static', static_url_path='')
app.config['MAX_CONTENT_LENGTH'] = int(os.getenv('MAX_REQUEST_MB', '220')) * 1024 * 1024
MAX_FILE_BYTES = int(os.getenv('MAX_FILE_MB', '100')) * 1024 * 1024
ALLOWED_TOOLS = {'merge','split','compress','pdf-word','pdf-ppt','word-pdf','ppt-pdf','pdf-jpg','jpg-pdf','pdf-excel','pdf-png','extract-images','flatten','pdf-info','excel-pdf','edit','sign','watermark','rotate','organize','delete','extract','page-numbers','protect','unlock','metadata','repair','ocr','crop','add-image','compare'}

@app.get('/')
def index(): return app.send_static_file('index.html')
@app.get('/health')
def health(): return jsonify(status='ok', service='AZ PDF LOVER API')

def pages_arg(value, count):
    if not value or not value.strip(): return list(range(count))
    out=[]
    for part in value.split(','):
        part=part.strip()
        if re.fullmatch(r'\d+', part): a=b=int(part)
        else:
            m=re.fullmatch(r'(\d+)\s*-\s*(\d+)', part)
            if not m: raise ValueError(f'Invalid page range: {part}')
            a,b=map(int,m.groups())
        if a < 1 or b < a or b > count: raise ValueError(f'Page range {part} is outside 1–{count}.')
        out.extend(range(a-1,b))
    return list(dict.fromkeys(out))

def pdf_open(path, password=''):
    doc=fitz.open(path)
    if doc.needs_pass:
        if not password or not doc.authenticate(password):
            doc.close(); raise ValueError('This PDF is password-protected. Enter the correct password if you are authorized to open it.')
    return doc

def office_convert(src, outdir, fmt='pdf'):
    exe=shutil.which('libreoffice') or shutil.which('soffice')
    if not exe: raise ValueError('Office conversion engine is unavailable. The provided Docker deployment includes LibreOffice; redeploy using the included Dockerfile.')
    p=subprocess.run([exe,'-env:UserInstallation=file://' + os.path.join(outdir,'lo-profile'), '--headless','--convert-to',fmt,'--outdir',outdir,src],capture_output=True,text=True,timeout=180)
    expected=os.path.join(outdir,Path(src).stem+'.'+fmt)
    if p.returncode or not os.path.exists(expected):
        raise ValueError('Office conversion failed. Check that the uploaded document is valid and not password-protected.')
    return expected

def text_from_pdf(doc): return '\n\n'.join(page.get_text('text') for page in doc)
def cleanbase(name): return Path(secure_filename(name or 'document')).stem[:80] or 'document'
def output_response(path, filename=None, mimetype=None):
    return send_file(path, as_attachment=True, download_name=filename or os.path.basename(path), mimetype=mimetype or 'application/octet-stream', max_age=0)

def get_files():
    files=request.files.getlist('files')
    if not files or all(not f.filename for f in files): raise ValueError('Please select the required file(s).')
    saved=[]
    for f in files:
        if not f.filename: continue
        name=secure_filename(f.filename)
        if not name: raise ValueError('Invalid filename.')
        data=f.read(MAX_FILE_BYTES+1)
        if len(data)>MAX_FILE_BYTES: raise ValueError(f'{name} exceeds the {MAX_FILE_BYTES//(1024*1024)} MB per-file limit.')
        path=os.path.join(request.environ['az_workdir'], f'{len(saved)}_{name}')
        with open(path,'wb') as out: out.write(data)
        saved.append((path,name))
    return saved

@app.post('/api/process')
def process():
    work=tempfile.mkdtemp(prefix='azpdf_')
    request.environ['az_workdir']=work
    try:
        tool=request.form.get('tool','').strip()
        if tool not in ALLOWED_TOOLS: raise ValueError('Unknown tool requested.')
        opts=json.loads(request.form.get('options','{}'))
        files=get_files(); paths=[p for p,n in files]; names=[n for p,n in files]
        out=os.path.join(work,'result')
        base=cleanbase(names[0])
        password=str(opts.get('password',''))
        pdf_tools={'split','compress','pdf-word','pdf-ppt','pdf-jpg','pdf-png','pdf-excel','extract-images','flatten','pdf-info','edit','sign','watermark','rotate','organize','delete','extract','page-numbers','protect','unlock','metadata','repair','ocr','crop','add-image','compare'}
        if tool in pdf_tools:
            if tool in {'add-image','sign'}:
                for p,n in files:
                    if n.lower().endswith('.pdf') or (tool=='sign' and n.lower().endswith(('.png','.jpg','.jpeg'))) or (tool=='add-image' and n.lower().endswith(('.png','.jpg','.jpeg'))): continue
                    raise ValueError(f'{n} must be a PDF'+(' or JPG/PNG image.' if tool in {'add-image','sign'} else ' file.'))
            else:
                for p,n in files:
                    if not n.lower().endswith('.pdf'): raise ValueError(f'{n} is not a PDF file.')
        if tool=='merge':
            if len(paths)<2: raise ValueError('Select at least two PDF files.')
            dst=fitz.open()
            for p,n in files:
                src=pdf_open(p,password); dst.insert_pdf(src); src.close()
            out += '.pdf'; dst.save(out,garbage=4,deflate=True); dst.close(); filename='merged.pdf'
        elif tool in {'split','extract','delete','organize','rotate','metadata','repair','crop','watermark','page-numbers','edit','sign','add-image','compress','protect','unlock'}:
            if tool=='sign' and any(n.lower().endswith(('.png','.jpg','.jpeg')) for n in names):
                pdf_entry=next(((p,n) for p,n in files if n.lower().endswith('.pdf')),None)
                image_entry=next(((p,n) for p,n in files if n.lower().endswith(('.png','.jpg','.jpeg'))),None)
                if not pdf_entry or not image_entry: raise ValueError('Choose a PDF and an image signature (PNG/JPG), or choose only a PDF to apply a text signature.')
                doc=pdf_open(pdf_entry[0],password); inds=pages_arg(opts.get('pages',''),len(doc)); x=float(opts.get('x') or 50); y=float(opts.get('y') or 50); width=float(opts.get('imageWidth') or 120)
                for i in inds:
                    page=doc[i]; ph=page.rect.height; page.insert_image(fitz.Rect(x,max(0,ph-y-width*.4),x+width,max(0,ph-y-width*.4)+width*.4),filename=image_entry[0],keep_proportion=True,overlay=True)
                out += '.pdf'; doc.save(out,garbage=4,deflate=True); doc.close(); filename=cleanbase(pdf_entry[1])+'_signed.pdf'
            elif tool=='add-image':
                if len(files)<2: raise ValueError('Select one PDF and one PNG/JPG image.')
                pdf_entry=next(((p,n) for p,n in files if n.lower().endswith('.pdf')),None)
                image_entry=next(((p,n) for p,n in files if n.lower().endswith(('.png','.jpg','.jpeg'))),None)
                if not pdf_entry or not image_entry: raise ValueError('Select one PDF and one JPG/PNG image.')
                src_path,src_name=pdf_entry; img_path,img_name=image_entry; doc=pdf_open(src_path,password); inds=pages_arg(opts.get('pages',''),len(doc))
                for i in inds:
                    page=doc[i]; rect=page.rect; width=float(opts.get('imageWidth') or 120); x=float(opts.get('x') or 50); y=float(opts.get('y') or 50)
                    page.insert_image(fitz.Rect(x, max(0,rect.height-y-width), x+width, max(0,rect.height-y-width)+width), filename=img_path, keep_proportion=True, overlay=True)
                out += '.pdf'; doc.save(out,garbage=4,deflate=True); doc.close(); filename=base+'_image.pdf'
            else:
                doc=pdf_open(paths[0],password)
                if tool=='split' or tool=='extract':
                    inds=pages_arg(opts.get('pages',''),len(doc)); dst=fitz.open(); dst.insert_pdf(doc,from_page=inds[0],to_page=inds[0]) if len(inds)==1 else None
                    if len(inds)!=1:
                        for i in inds: dst.insert_pdf(doc,from_page=i,to_page=i)
                    out += '.pdf'; dst.save(out,garbage=4,deflate=True); dst.close(); filename=base+'_pages.pdf'
                elif tool=='delete':
                    inds=pages_arg(opts.get('pages',''),len(doc))
                    if len(inds)>=len(doc): raise ValueError('You cannot delete every page; keep at least one page.')
                    for i in sorted(inds,reverse=True): doc.delete_page(i)
                    out += '.pdf'; doc.save(out,garbage=4,deflate=True); filename=base+'_deleted.pdf'
                elif tool=='organize':
                    raw=opts.get('pages','').strip()
                    if not raw: raise ValueError('Enter the new page order, e.g. 3,1,2.')
                    order=pages_arg(raw,len(doc))
                    if len(order)!=len(doc) or len(set(order))!=len(doc): raise ValueError(f'Enter each page exactly once, e.g. {",".join(map(str,range(len(doc),0,-1)))}.')
                    dst=fitz.open()
                    for i in order: dst.insert_pdf(doc,from_page=i,to_page=i)
                    out += '.pdf'; dst.save(out,garbage=4,deflate=True); dst.close(); filename=base+'_organized.pdf'
                elif tool=='rotate':
                    deg=int(opts.get('degrees') or 90)
                    for i in pages_arg(opts.get('pages',''),len(doc)): doc[i].set_rotation((doc[i].rotation+deg)%360)
                    out += '.pdf'; doc.save(out,garbage=4,deflate=True); filename=base+'_rotated.pdf'
                elif tool=='metadata':
                    meta=doc.metadata or {}; meta.update({'title':opts.get('title',''),'author':opts.get('author',''),'subject':opts.get('subject',''),'keywords':opts.get('keywords','')}); doc.set_metadata(meta)
                    out += '.pdf'; doc.save(out,garbage=4,deflate=True); filename=base+'_metadata.pdf'
                elif tool=='repair':
                    doc.close(); doc=fitz.open(paths[0]); out += '.pdf'; doc.save(out,garbage=4,deflate=True,clean=True); doc.close(); filename=base+'_repaired.pdf'
                elif tool=='crop':
                    x=float(opts.get('left') or 0); y=float(opts.get('bottom') or 0); w=float(opts.get('width') or 500); h=float(opts.get('height') or 700)
                    if w<=0 or h<=0: raise ValueError('Crop width and height must be positive.')
                    for page in doc:
                        r=page.rect; # Convert user bottom-origin PDF points into MuPDF top-origin coordinates
                        box=fitz.Rect(x,max(0,r.height-y-h),min(r.width,x+w),min(r.height,max(0,r.height-y-h)+h))
                        if box.is_empty or box.width<=0 or box.height<=0: raise ValueError('Crop box is outside a page boundary.')
                        page.set_cropbox(box)
                    out += '.pdf'; doc.save(out,garbage=4,deflate=True); filename=base+'_cropped.pdf'
                elif tool=='watermark':
                    text=str(opts.get('watermarkText') or 'CONFIDENTIAL')[:300]; opacity=max(0,min(1,float(opts.get('opacity') or .25))); fontsize=max(4,min(144,float(opts.get('fontSize') or 36)))
                    for page in doc:
                        r=page.rect; pos=opts.get('position','center'); point=(r.width*.12,r.height*.5) if pos=='center' else (r.width*.1,50 if pos=='top' else r.height-30)
                        page.insert_text(point,text,fontsize=fontsize,color=(.45,.45,.45),overlay=True,rotate=0,fill_opacity=opacity)
                    out += '.pdf'; doc.save(out,garbage=4,deflate=True); filename=base+'_watermarked.pdf'
                elif tool=='page-numbers':
                    start=int(opts.get('startNumber') or 1); pos=opts.get('position','bottom-center')
                    for i,page in enumerate(doc):
                        r=page.rect; text=str(start+i); y=(page.rect.height-25) if pos!='top-right' else 35; x=r.width-55 if pos in ('bottom-right','top-right') else r.width/2-10
                        page.insert_text((x,y),text,fontsize=11,color=(.25,.25,.25),overlay=True)
                    out += '.pdf'; doc.save(out,garbage=4,deflate=True); filename=base+'_numbered.pdf'
                elif tool in {'edit','sign'}:
                    text=str(opts.get('overlayText') or ('Signed' if tool=='sign' else ''))[:500]
                    if text:
                        for i in pages_arg(opts.get('pages',''),len(doc)):
                            page=doc[i]; y=page.rect.height-float(opts.get('y') or 50); x=float(opts.get('x') or 50)
                            page.insert_text((x,y),text,fontsize=22 if tool=='sign' else 14,color=(.08,.1,.18),overlay=True)
                    out += '.pdf'; doc.save(out,garbage=4,deflate=True); filename=base+('_signed.pdf' if tool=='sign' else '_edited.pdf')
                elif tool=='compress':
                    out += '.pdf'; doc.save(out,garbage=4,deflate=True,clean=True,deflate_images=True,deflate_fonts=True); filename=base+'_compressed.pdf'
                    # Optional stronger raster/image optimization via Ghostscript if installed
                    gs=shutil.which('gs')
                    if gs:
                        tmp=os.path.join(work,'gs.pdf'); preset=opts.get('compressMode','ebook'); preset=preset if preset in ('screen','ebook','printer','prepress','default') else 'ebook'
                        doc.save(out,garbage=4,deflate=True,clean=True); doc.close()
                        p=subprocess.run([gs,'-sDEVICE=pdfwrite','-dCompatibilityLevel=1.5','-dNOPAUSE','-dBATCH','-dQUIET',f'-dPDFSETTINGS=/{preset}',f'-sOutputFile={tmp}',out],capture_output=True,timeout=180)
                        if p.returncode==0 and os.path.exists(tmp) and os.path.getsize(tmp)<os.path.getsize(out): shutil.copy2(tmp,out)
                        return output_response(out,filename,'application/pdf')
                elif tool=='protect':
                    doc.close()
                    if not password: raise ValueError('Enter a password to protect the PDF.')
                    with pikepdf.open(paths[0]) as pdf:
                        out += '.pdf'; pdf.save(out,encryption=pikepdf.Encryption(owner=password,user=password,R=6,allow=pikepdf.Permissions(extract=False,modify_annotation=False,modify_assembly=False,modify_form=False,modify_other=False,print_lowres=True,print_highres=False)))
                    filename=base+'_protected.pdf'
                elif tool=='unlock':
                    doc.close()
                    if not password: raise ValueError('Enter the current PDF password. This tool only works for documents you are authorized to access.')
                    with pikepdf.open(paths[0],password=password) as pdf: out += '.pdf'; pdf.save(out)
                    filename=base+'_unlocked.pdf'
                doc.close() if tool not in ('repair','protect','unlock','compress') else None
        elif tool=='pdf-word':
            doc=pdf_open(paths[0],password); txt=text_from_pdf(doc); doc.close()
            if opts.get('wordMode','docx')=='txt': out += '.txt'; Path(out).write_text(txt,encoding='utf-8'); filename=base+'.txt'
            else:
                d=Document()
                for line in txt.splitlines(): d.add_paragraph(line)
                out += '.docx'; d.save(out); filename=base+'.docx'
        elif tool=='pdf-ppt':
            doc=pdf_open(paths[0],password); prs=Presentation(); layout=opts.get('slideSize','LAYOUT_WIDE'); prs.slide_width=PptInches(13.333 if layout!='LAYOUT_4X3' else 10); prs.slide_height=PptInches(7.5)
            for page in doc:
                pix=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False); img=os.path.join(work,f'p{page.number}.png'); pix.save(img); slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.shapes.add_picture(img,0,0,width=prs.slide_width,height=prs.slide_height)
            doc.close(); out += '.pptx'; prs.save(out); filename=base+'.pptx'
        elif tool=='word-pdf' or tool=='ppt-pdf' or tool=='excel-pdf':
            src=paths[0]; ext=Path(names[0]).suffix.lower()
            if tool=='word-pdf' and ext not in ('.docx','.doc'): raise ValueError('Upload a Word DOCX or DOC file.')
            if tool=='ppt-pdf' and ext not in ('.pptx','.ppt'): raise ValueError('Upload a PowerPoint PPTX or PPT file.')
            if tool=='excel-pdf' and ext not in ('.xlsx','.xls'): raise ValueError('Upload an Excel XLSX or XLS file.')
            if tool=='excel-pdf' and ext=='.xlsx':
                # Export a selected XLSX worksheet so LibreOffice prints only the requested sheet.
                wb=load_workbook(src,read_only=False,data_only=True)
                sheetname=opts.get('sheetName') or wb.sheetnames[0]
                if sheetname not in wb.sheetnames: sheetname=wb.sheetnames[0]
                for ws in list(wb.worksheets):
                    if ws.title!=sheetname: wb.remove(ws)
                tmpxlsx=os.path.join(work,'worksheet.xlsx'); wb.save(tmpxlsx); src=tmpxlsx
            outpdf=office_convert(src,work,'pdf'); out=outpdf; filename=base+'.pdf'
        elif tool=='jpg-pdf':
            if not paths: raise ValueError('Select one or more JPG images.')
            dst=fitz.open(); size=opts.get('pageSize','image')
            for p,n in files:
                if Path(n).suffix.lower() not in ('.jpg','.jpeg'): raise ValueError('JPG to PDF accepts JPG/JPEG files only.')
                pix=fitz.Pixmap(p); w,h=pix.width,pix.height
                if size=='a4': rect=fitz.Rect(0,0,595.28,841.89)
                elif size=='letter': rect=fitz.Rect(0,0,612,792)
                else: rect=fitz.Rect(0,0,w,h)
                page=dst.new_page(width=rect.width,height=rect.height); scale=min(rect.width/w,rect.height/h); rw,rh=w*scale,h*scale; page.insert_image(fitz.Rect((rect.width-rw)/2,(rect.height-rh)/2,(rect.width+rw)/2,(rect.height+rh)/2),filename=p)
            out += '.pdf'; dst.save(out,garbage=4,deflate=True); dst.close(); filename='images.pdf'
        elif tool=='pdf-jpg':
            doc=pdf_open(paths[0],password); zpath=out+'.zip'; scale=max(.5,min(4,float(opts.get('scale') or 1.5))); quality=max(30,min(100,int(float(opts.get('quality') or .9)*100)))
            with zipfile.ZipFile(zpath,'w',zipfile.ZIP_DEFLATED) as z:
                for i,page in enumerate(doc):
                    pix=page.get_pixmap(matrix=fitz.Matrix(scale,scale),alpha=False); im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples); b=BytesIO(); im.save(b,format='JPEG',quality=quality); z.writestr(f'page-{i+1:03}.jpg',b.getvalue())
            doc.close(); out=zpath; filename=base+'_jpg_pages.zip'
        elif tool=='pdf-excel':
            doc=pdf_open(paths[0],password); text=text_from_pdf(doc); doc.close()
            if opts.get('excelMode')=='csv': out += '.csv'; Path(out).write_text('\n'.join('"'+x.replace('"','""')+'"' for x in text.splitlines()),encoding='utf-8-sig'); filename=base+'.csv'
            else:
                wb=Workbook(); ws=wb.active; ws.title='Extracted text'
                for line in text.splitlines(): ws.append([line])
                out += '.xlsx'; wb.save(out); filename=base+'.xlsx'
        elif tool=='ocr':
            doc=pdf_open(paths[0],password); langs={'eng':'eng','urd':'urd','spa':'spa','fra':'fra','deu':'deu'}; lang=langs.get(opts.get('language','eng'),'eng'); inds=pages_arg(opts.get('pages',''),len(doc)); chunks=[]
            for i in inds:
                pix=doc[i].get_pixmap(matrix=fitz.Matrix(2,2),alpha=False); im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
                try: recognized=pytesseract.image_to_string(im,lang=lang)
                except Exception as e: raise ValueError(f'OCR language {lang} is unavailable on this server. Check the Docker language-pack installation.') from e
                chunks.append(f'--- Page {i+1} ---\n{recognized}')
            doc.close(); out += '.txt'; Path(out).write_text('\n\n'.join(chunks),encoding='utf-8'); filename=base+'_ocr.txt'
        elif tool=='pdf-png':
            doc=pdf_open(paths[0],password); zpath=out+'.zip'; scale=max(.5,min(4,float(opts.get('scale') or 1.5)))
            with zipfile.ZipFile(zpath,'w',zipfile.ZIP_DEFLATED) as z:
                for i,page in enumerate(doc):
                    pix=page.get_pixmap(matrix=fitz.Matrix(scale,scale),alpha=False); z.writestr(f'page-{i+1:03}.png',pix.tobytes('png'))
            doc.close(); out=zpath; filename=base+'_png_pages.zip'
        elif tool=='extract-images':
            doc=pdf_open(paths[0],password); zpath=out+'.zip'; seen=set(); count=0
            with zipfile.ZipFile(zpath,'w',zipfile.ZIP_DEFLATED) as z:
                for page in doc:
                    for image in page.get_images(full=True):
                        xref=image[0]
                        if xref in seen: continue
                        seen.add(xref); item=doc.extract_image(xref); ext=item.get('ext','png'); z.writestr(f'image-{count+1:03}.{ext}',item['image']); count+=1
            doc.close()
            if count==0: raise ValueError('No embedded raster images were found in this PDF.')
            out=zpath; filename=base+'_extracted_images.zip'
        elif tool=='flatten':
            doc=pdf_open(paths[0],password); flat=fitz.open()
            for page in doc:
                pix=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False); target=flat.new_page(width=page.rect.width,height=page.rect.height); target.insert_image(target.rect,stream=pix.tobytes('png'))
            doc.close(); out += '.pdf'; flat.save(out,garbage=4,deflate=True); flat.close(); filename=base+'_flattened.pdf'
        elif tool=='pdf-info':
            doc=pdf_open(paths[0],password); meta=doc.metadata or {}; lines=[f'File: {names[0]}',f'Pages: {len(doc)}',f'PDF version: {doc.metadata.get("format", "")}',f'Encrypted: {bool(doc.is_encrypted)}','', 'Metadata:']
            for key in ('title','author','subject','keywords','creator','producer','creationDate','modDate'): lines.append(f'{key}: {meta.get(key, "")}')
            lines += ['', 'Pages:']
            for i,page in enumerate(doc): lines.append(f'{i+1}: {page.rect.width:.1f} × {page.rect.height:.1f} pt; rotation={page.rotation}°; text characters={len(page.get_text("text"))}')
            doc.close(); out += '.txt'; Path(out).write_text('\n'.join(lines),encoding='utf-8'); filename=base+'_info.txt'
        elif tool=='compare':
            if len(paths)!=2: raise ValueError('Choose exactly two PDF files to compare.')
            a=pdf_open(paths[0],password); b=pdf_open(paths[1],password); ta=text_from_pdf(a); tb=text_from_pdf(b)
            import difflib
            report=f'AZ PDF LOVER — PDF comparison\n\nFile A: {names[0]} ({len(a)} pages)\nFile B: {names[1]} ({len(b)} pages)\n\n'+'\n'.join(difflib.unified_diff(ta.splitlines(),tb.splitlines(),fromfile=names[0],tofile=names[1],lineterm=''))
            a.close(); b.close(); out += '.txt'; Path(out).write_text(report,encoding='utf-8'); filename='comparison_'+base+'.txt'
        else: raise ValueError('This tool is not available.')
        if not os.path.isfile(out): raise ValueError('Processing did not produce an output file. Please try another valid document.')
        return output_response(out,filename)
    except ValueError as e:
        return jsonify(error=str(e)),400
    except subprocess.TimeoutExpired:
        return jsonify(error='This document took too long to process. Try a smaller file or fewer pages.'),408
    except Exception as e:
        app.logger.exception('PDF processing failed')
        return jsonify(error='Could not process this file. It may be damaged, encrypted, unsupported, or too complex. Please check the file and options.'),400
    finally:
        shutil.rmtree(work,ignore_errors=True)

@app.errorhandler(413)
def too_large(_): return jsonify(error='Upload too large. Maximum total request size is 220 MB; each file may be up to 100 MB.'),413

if __name__=='__main__': app.run(host='0.0.0.0',port=int(os.getenv('PORT','10000')))
